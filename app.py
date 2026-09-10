#!/usr/bin/env python3
"""Santa Cruz 2027 - Montagem de Elenco.

App proprio, porta 5090. Nao faz parte do hub Botafogo (:5555).
Base de jogadores: dados/jogadores.json (gerada por preparar_base.py).
Cenarios salvos: dados/cenarios.json.
"""
import io
import json
import os
import uuid
from datetime import datetime

from flask import Flask, Response, jsonify, request, render_template, send_file, abort

AQUI = os.path.dirname(os.path.abspath(__file__))
ARQ_JOGADORES = os.path.join(AQUI, "dados", "jogadores.json")
ARQ_CENARIOS = os.environ.get("SC_CENARIOS") or os.path.join(AQUI, "dados", "cenarios.json")
ARQ_PREMISSAS = os.environ.get("SC_PREMISSAS") or os.path.join(AQUI, "dados", "premissas.json")

# As regras do projeto, ditas pelo usuario ao longo da montagem. Ficam no servidor e
# nao no navegador de proposito: sao o combinado do trabalho, nao preferencia de tela.
# Este e so o conteudo INICIAL — a partir da primeira gravacao vale o arquivo.
PREMISSAS_INICIAIS = [
    ("Montagem do elenco", "Time físico",
     "Elenco montado para correr: a intensidade é critério de escolha, não detalhe. "
     "É por isso que a aba Físico compara todo candidato com a régua das Séries A e B."),
    ("Montagem do elenco", "Muitos minutos por temporada",
     "Priorizar quem joga, e joga muito. Minutagem alta e repetida nas últimas três "
     "temporadas vale mais do que um pico isolado."),
    ("Montagem do elenco", "Titulares consolidados, reservas com potencial",
     "Os onze são gente pronta, com rodagem comprovada. O banco é onde entra a apost"
     "a: jovem com potencial de crescer dentro da temporada."),
    ("Montagem do elenco", "Goleiro top — investir",
     "Posição em que vale pagar acima da média do elenco. Goleiro decide pontos na S"
     "érie B e o custo de errar aqui é alto."),
    ("Montagem do elenco", "Comissão técnica top — investir",
     "A comissão entra no teto de R$ 2,8 MM e é para ser boa, não barata. O que se g"
     "asta nela sai da massa salarial de propósito."),
    ("Montagem do elenco", "Salário baixo, premiação alta por vitória e acesso",
     "Contratar abaixo do que o mercado pagaria e pendurar o dinheiro grande na "
     "premiação: bicho por vitória ao longo do campeonato e um prêmio forte pelo "
     "acesso. O risco vai para o resultado — ganhando, paga-se muito; não ganhando, "
     "o custo fixo não afunda o clube. A premiação é variável e não entra no teto "
     "mensal de R$ 2,8 MM, que é custo recorrente."),
    ("Montagem do elenco", "Logística diferenciada para a Série B",
     "A Série B tem viagem longa e calendário apertado. Estrutura de deslocamento, d"
     "escanso e recuperação entra na conta da montagem, não é despesa à parte."),

    ("Orçamento", "Teto de custo total",
     "R$ 2.800.000 por mês é o custo total máximo do elenco."),
    ("Orçamento", "O teto já inclui a comissão técnica",
     "A comissão entra dentro dos R$ 2,8 MM, não por fora. Padrão de R$ 300.000, "
     "editável ou detalhada por cargo no modal Orçamento."),
    ("Orçamento", "Encargos de 1,25×",
     "O custo real de um jogador é 1,25 vez o salário oferecido a ele. "
     "Massa salarial disponível = (teto − comissão) ÷ 1,25."),
    ("Orçamento", "O salário digitado é o do jogador",
     "O número no card é o que o atleta recebe, sem encargos. Quem aplica o 1,25 é a conta."),
    ("Orçamento", "Sem salário sugerido",
     "Jogador entra no campograma com salário 0. A faixa do TransferRoom aparece na "
     "ficha e no Fim de contrato, mas não preenche o card."),
    ("Orçamento", "Cotação do euro",
     "A faixa salarial do TransferRoom é anual em euros; vira mensal em reais pela "
     "cotação do modal Orçamento (padrão 6,30)."),

    ("Elenco", "Objetivo",
     "Montar o elenco de 2027 do Santa Cruz visando o acesso à Série B."),
    ("Elenco", "Onze posições, três opções",
     "GK, LB, LCB, RCB, RB, DM, CM, AM, LW, CF, RW — com 3 vagas por posição, "
     "ajustáveis uma a uma."),
    ("Elenco", "Limite de estrangeiros",
     "9 estrangeiros no elenco."),
    ("Elenco", "Siglas internacionais na tela",
     "GK, LCB, RCB, LB, RB, DM, CM, AM, LW, RW, CF. A chave interna continua em "
     "português porque indexa a base, os KPIs e os cenários já gravados."),

    ("Dados", "O cadastro vem do fim de contrato",
     "40.059 jogadores do levantamento de fim de contrato (Transfermarkt), não do "
     "ranking — o ranking só enxerga quem tem minutagem (18.281)."),
    ("Dados", "Capology não cobre o Brasil",
     "Não publica salário do futebol brasileiro. Todo salário é digitado, exceto a "
     "estimativa do TransferRoom."),
    ("Dados", "Histórico de três temporadas",
     "2024, 2025 e 2026, dos mesmos Excels do Wyscout que alimentam o ranking."),
    ("Dados", "Ligação entre temporadas",
     "Primeiro por id (player_uid). Por nome só quando o nome é único na temporada de "
     "hoje e na de lá, conferindo país, altura e idade. Nome repetido não casa: o "
     "Pedro do Flamengo chegou a receber a temporada de um Pedro do Guabirá."),
    ("Dados", "O Wyscout não marca a origem do gol",
     "Não existe \"gol de bola parada\" na base. O que existe são os dois lados: quem "
     "cobra (escanteios + faltas por 90, pênaltis) e quem cabeceia (gols de cabeça)."),
    ("Dados", "Goleador recorrente",
     "Quem fez 5 gols ou mais em pelo menos 2 das 3 temporadas."),
    ("Dados", "O oGol tem jogos, não minutos",
     "Onde não houver minutos, mostrar jogos, marcado na tela. Nada é estimado."),

    ("Físico", "A régua é a Série A + B",
     "Todo percentil é calculado entre os jogadores da mesma posição com tracking nas "
     "Séries A e B. Jogadores de outras ligas entram nessa mesma régua para poder comparar."),
    ("Físico", "Nos tempos, menor é melhor",
     "Nos indicadores de tempo (até o sprint, 505, após girar), liderar é ter o número "
     "menor. A cor não segue o sinal da diferença."),
    ("Físico", "Índice físico",
     "Índice do grupo = média dos percentis do grupo. Índice geral = média dos cinco "
     "grupos. É ele que escolhe os 5 melhores de cada série."),
    ("Físico", "Contorno diz o nível",
     "Verde quando o jogador bate as médias das duas séries; âmbar quando bate só a "
     "mais fraca das duas."),
    ("Físico", "Estudo Série A × Série B",
     "Só entram posições com pelo menos três jogadores com tracking de cada lado — o "
     "que deixa o goleiro de fora."),
    ("Físico", "Comparativo gravado é lista, não receita",
     "O que se grava são os atletas que estavam na tela, com os incluídos e os "
     "excluídos. Guardar \"top 5 + filtro\" faria o comparativo mudar sozinho quando "
     "a base fosse regerada."),
]


def _premissas_padrao():
    return [{"id": f"p{i + 1}", "grupo": g, "titulo": t, "texto": x, "fonte": "inicial"}
            for i, (g, t, x) in enumerate(PREMISSAS_INICIAIS)]


def ler_premissas():
    if not os.path.exists(ARQ_PREMISSAS):
        return _premissas_padrao()
    try:
        with open(ARQ_PREMISSAS, encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return _premissas_padrao()


def gravar_premissas(lista):
    os.makedirs(os.path.dirname(ARQ_PREMISSAS), exist_ok=True)
    tmp = ARQ_PREMISSAS + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(lista, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, ARQ_PREMISSAS)
PORTA = 5090

app = Flask(__name__)
app.json.ensure_ascii = False
# app local: template e static sempre frescos, sem cache do navegador
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0
app.jinja_env.auto_reload = True

def senha_exigida():
    """Senha de acesso: definida por variavel de ambiente na hospedagem.
    Sem ela (rodando local), o app abre direto."""
    return os.environ.get("SC_SENHA", "").strip()


@app.before_request
def exigir_senha():
    senha = senha_exigida()
    if not senha:
        return None
    auth = request.authorization
    if auth and auth.password == senha:
        return None
    return Response(
        "Acesso restrito ao Santa Cruz 2027.", 401,
        {"WWW-Authenticate": 'Basic realm="Santa Cruz 2027"'})


@app.route("/dados/<path:arquivo>")
def dados(arquivo):
    """Bases servidas como arquivo: nada fica carregado na memoria do servidor."""
    caminho = os.path.normpath(os.path.join(AQUI, "dados", arquivo))
    if not caminho.startswith(os.path.join(AQUI, "dados")) or not os.path.exists(caminho):
        abort(404)
    if os.path.basename(caminho) == "cenarios.json":   # nao e publico
        abort(404)
    return send_file(caminho, max_age=3600)


def ler_cenarios():
    if not os.path.exists(ARQ_CENARIOS):
        return {}
    try:
        with open(ARQ_CENARIOS, encoding="utf-8") as fh:
            return json.load(fh)
    except (json.JSONDecodeError, OSError):
        return {}


def gravar_cenarios(dados):
    os.makedirs(os.path.dirname(ARQ_CENARIOS), exist_ok=True)
    tmp = ARQ_CENARIOS + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(dados, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, ARQ_CENARIOS)


def versao_dados():
    """Assinatura das bases, para o navegador nao ficar com dados velhos."""
    marcas = []
    # todos os arquivos que o app.js busca em /dados: faltar um aqui significa navegador
    # servindo base velha em silencio — foi o que aconteceu com o raio_ref.json
    for nome in ("jogadores.json", "historico.json", "premissas.json",
                 "raio_ref.json", "posicao_overrides.json"):
        caminho = os.path.join(AQUI, "dados", nome)
        marcas.append(str(int(os.path.getmtime(caminho))) if os.path.exists(caminho) else "0")
    return "-".join(marcas)


def versao_estatica():
    """Assinatura dos arquivos de front, para o navegador nunca servir versao velha."""
    marcas = []
    for nome in ("app.js", "style.css", "fs_visoes.js"):
        caminho = os.path.join(AQUI, "static", nome)
        marcas.append(str(int(os.path.getmtime(caminho))) if os.path.exists(caminho) else "0")
    return "-".join(marcas)


@app.route("/")
def index():
    resp = app.make_response(render_template("index.html", ver=versao_estatica(),
                                             verDados=versao_dados()))
    resp.headers["Cache-Control"] = "no-store"
    return resp


@app.route("/escudo")
def escudo():
    """Serve o escudo que o usuario colocar em static/ (escudo.png/.svg/.jpg/.webp).
    Sem arquivo, devolve 404 e a tela cai no brasao desenhado."""
    for nome in ("escudo.svg", "escudo.png", "escudo.jpg", "escudo.jpeg", "escudo.webp"):
        caminho = os.path.join(AQUI, "static", nome)
        if os.path.exists(caminho):
            return send_file(caminho, max_age=0)
    abort(404)


@app.route("/api/diagnostico", methods=["POST"])
def api_diagnostico():
    """Recebe o rastro de um tremor detectado na tela e grava para analise."""
    dados = request.get_json(silent=True) or {}
    caminho = os.path.join(AQUI, "dados", "diagnostico.log")
    with open(caminho, "a", encoding="utf-8") as fh:
        fh.write(datetime.now().strftime("%Y-%m-%d %H:%M:%S") + " " +
                 json.dumps(dados, ensure_ascii=False) + "\n")
    return jsonify({"ok": True})


@app.route("/api/comparativo")
def api_comparativo():
    """Resumo de cada grupo salvo, para a tela de comparacao."""
    SETOR = {"gol": ["GOL"], "defesa": ["LD", "LE", "ZD", "ZE"],
             "meio": ["VOL", "MED", "MEI"], "ataque": ["EE", "ED", "CA"]}
    saida = []
    for cid, c in ler_cenarios().items():
        elenco = c.get("elenco") or {}
        atletas = [j for lista in elenco.values() for j in lista]
        folha = sum(float(j.get("salario") or 0) for j in atletas)
        fator = float(c.get("fator") or 1) or 1
        comissao = float(c.get("comissao") or 0)
        teto = float(c.get("teto") or 0)
        idades = [float(j["idade"]) for j in atletas if j.get("idade")]
        setores = {}
        for s, poss in SETOR.items():
            setores[s] = sum(float(j.get("salario") or 0)
                             for p in poss for j in elenco.get(p, []))
        saida.append({
            "id": cid, "nome": c.get("nome", "sem nome"),
            "atualizado": c.get("atualizado", ""),
            "atletas": len(atletas),
            "estrangeiros": sum(1 for j in atletas if j.get("estrangeiro")),
            "folha": folha, "fator": fator, "comissao": comissao, "teto": teto,
            "custoElenco": folha * fator,
            "custoTotal": folha * fator + comissao,
            "sobra": teto - (folha * fator + comissao),
            "media": (folha / len(atletas)) if atletas else 0,
            "maior": max((float(j.get("salario") or 0) for j in atletas), default=0),
            "idadeMedia": (sum(idades) / len(idades)) if idades else 0,
            "semSalario": sum(1 for j in atletas if not j.get("salario")),
            "setores": setores,
        })
    saida.sort(key=lambda x: x["atualizado"], reverse=True)
    return jsonify(saida)


@app.route("/api/premissas", methods=["GET", "POST"])
def api_premissas():
    """As premissas do projeto. GET devolve a lista; POST grava a lista inteira."""
    if request.method == "GET":
        return jsonify(ler_premissas())
    lista = request.get_json(silent=True)
    if not isinstance(lista, list):
        return jsonify({"erro": "esperava uma lista"}), 400
    limpa = []
    for p in lista[:400]:
        if not isinstance(p, dict):
            continue
        limpa.append({
            "id": str(p.get("id") or "")[:40],
            "grupo": str(p.get("grupo") or "Geral")[:60],
            "titulo": str(p.get("titulo") or "")[:140],
            "texto": str(p.get("texto") or "")[:2000],
            "fonte": str(p.get("fonte") or "usuario")[:20],
        })
    gravar_premissas(limpa)
    return jsonify({"ok": True, "total": len(limpa)})


@app.route("/api/cenarios", methods=["GET", "POST"])
def api_cenarios():
    cen = ler_cenarios()
    if request.method == "GET":
        return jsonify([
            {"id": k, "nome": v.get("nome", "sem nome"),
             "atualizado": v.get("atualizado", ""),
             "total": v.get("total", 0), "atletas": v.get("atletas", 0)}
            for k, v in sorted(cen.items(), key=lambda kv: kv[1].get("atualizado", ""), reverse=True)
        ])

    corpo = request.get_json(silent=True) or {}
    cid = corpo.get("id") or uuid.uuid4().hex[:10]
    corpo["id"] = cid
    corpo["atualizado"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    cen[cid] = corpo
    gravar_cenarios(cen)
    return jsonify({"ok": True, "id": cid, "atualizado": corpo["atualizado"]})


@app.route("/api/cenario/<cid>", methods=["GET", "DELETE"])
def api_cenario(cid):
    cen = ler_cenarios()
    if cid not in cen:
        abort(404)
    if request.method == "DELETE":
        cen.pop(cid)
        gravar_cenarios(cen)
        return jsonify({"ok": True})
    return jsonify(cen[cid])


@app.route("/api/exportar", methods=["POST"])
def api_exportar():
    """Recebe o elenco montado e devolve um .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    dados = request.get_json(silent=True) or {}
    nome_cen = dados.get("nome") or "Elenco Santa Cruz 2027"
    teto = float(dados.get("teto") or 0)
    comissao = float(dados.get("comissao") or 0)
    fator = float(dados.get("fator") or 1) or 1
    disponivel = float(dados.get("disponivel") or ((teto - comissao) / fator))
    linhas = dados.get("linhas") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "Elenco 2027"

    tit = Font(bold=True, color="FFFFFF", size=11)
    fundo = PatternFill("solid", fgColor="8B0000")
    negrito = Font(bold=True)

    ws.append([nome_cen])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Custo total maximo (mes)", teto])
    ws.append(["Comissao tecnica (mes)", comissao])
    ws.append(["Fator de encargos", fator])
    ws.append(["Massa salarial disponivel", disponivel])
    for r in (2, 3, 5):
        ws.cell(row=r, column=2).number_format = 'R$ #,##0'
    ws.cell(row=4, column=2).number_format = '0.00'
    ws.append([])

    cab = ["Posicao", "Jogador", "Clube", "Liga", "Idade", "Overall", "Contrato",
           "Nacionalidade", "Estrangeiro", "Status",
           "Salario mensal (R$)", "Custo p/ o clube (R$)", "% da massa"]
    ws.append(cab)
    lin_cab = ws.max_row
    for c in range(1, len(cab) + 1):
        cel = ws.cell(row=lin_cab, column=c)
        cel.font = tit
        cel.fill = fundo
        cel.alignment = Alignment(horizontal="center")

    total = 0.0
    n_estr = 0
    for ln in linhas:
        sal = float(ln.get("salario") or 0)
        total += sal
        if str(ln.get("estrangeiro", "")).upper() == "SIM":
            n_estr += 1
        ws.append([
            ln.get("posicao", ""), ln.get("nome", ""), ln.get("clube", ""),
            ln.get("liga", ""), ln.get("idade", ""), ln.get("overall", ""),
            ln.get("contrato", ""), ln.get("nacionalidade", ""), ln.get("estrangeiro", ""),
            ln.get("status", ""), sal, sal * fator,
            (sal / disponivel) if disponivel else 0,
        ])

    ult = ws.max_row + 1
    ws.cell(row=ult, column=10, value="TOTAL").font = negrito
    ws.cell(row=ult, column=11, value=total).font = negrito
    ws.cell(row=ult, column=12, value=total * fator).font = negrito
    ws.cell(row=ult, column=13, value=(total / disponivel) if disponivel else 0).font = negrito

    ws.append([])
    for rot, val, fmt in (
        ("Atletas", len(linhas), "0"),
        ("Estrangeiros", n_estr, "0"),
        ("Folha salarial (jogadores)", total, 'R$ #,##0'),
        ("Custo do elenco (com encargos)", total * fator, 'R$ #,##0'),
        ("Custo total (elenco + comissao)", total * fator + comissao, 'R$ #,##0'),
        ("Sobra sobre o custo maximo", teto - (total * fator + comissao), 'R$ #,##0'),
    ):
        ws.append([rot, val])
        ws.cell(row=ws.max_row, column=1).font = negrito
        ws.cell(row=ws.max_row, column=2).number_format = fmt

    for r in range(lin_cab + 1, ult + 1):
        ws.cell(row=r, column=11).number_format = 'R$ #,##0'
        ws.cell(row=r, column=12).number_format = 'R$ #,##0'
        ws.cell(row=r, column=13).number_format = '0.0%'
    for c, larg in enumerate([18, 28, 20, 12, 7, 8, 12, 15, 11, 15, 19, 20, 11], start=1):
        ws.column_dimensions[get_column_letter(c)].width = larg
    ws.freeze_panes = ws.cell(row=lin_cab + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    seguro = "".join(ch for ch in nome_cen if ch.isalnum() or ch in " -_").strip() or "elenco"
    return send_file(buf, as_attachment=True,
                     download_name=f"{seguro}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    if not os.path.exists(ARQ_JOGADORES):
        raise SystemExit("dados/jogadores.json nao existe. Rode: python3 preparar_base.py ago26")
    porta = int(os.environ.get("PORT", PORTA))
    print(f"\n  Santa Cruz 2027 - Montagem de Elenco")
    print(f"  http://localhost:{porta}\n")
    app.run(host="0.0.0.0" if os.environ.get("PORT") else "127.0.0.1", port=porta, debug=False)
